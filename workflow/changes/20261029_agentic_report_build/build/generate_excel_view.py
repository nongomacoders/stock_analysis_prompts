"""Generate the consolidated static values-only Excel view from report.sqlite.

Governed by CHARTER covenants:
  - Derived tools read report.sqlite, never the PDF twice.
  - Covenant 2: All arithmetic executes in code. NO FORMULAS in any cell.
    The workbook contains static values only (strings, integers, floats).
  - Tabs:
      1. Primary_Statements (192 rows): Balance Sheet, Income Statement, Cash Flows (pp.17-18)
      2. Production_Costs_p48_50 (1,196 rows): Group Production & Costs Summary (pp.48-50)
      3. Conflicts (78 rows): Audited arithmetic & structural conflicts
      4. Questions (8 rows): Tracked ledger questions with status & lineage
"""
from __future__ import annotations

import pathlib
import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = pathlib.Path(__file__).resolve().parent.parent / "build" / "report.sqlite"
XLSX_PATH = pathlib.Path(__file__).resolve().parent.parent / "build" / "report_view.xlsx"


def build_excel_view():
    con = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill_blue = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_fill_green = PatternFill(start_color="274E13", end_color="274E13", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # ------------------------------------------------ 1. Primary Statements Tab
    ws_primary = wb.create_sheet(title="Primary_Statements")
    headers_primary = [
        "Cell ID", "Physical Page", "Statement Region", "Row Label", "Fiscal Year",
        "Value Raw", "Value Num (USD '000)", "Unit", "Extraction Commit"
    ]
    ws_primary.append(headers_primary)

    q_primary = """
        SELECT id, physical_page, table_region, row_label, fiscal_year,
               value_raw, value_num, unit, extraction_commit
        FROM cells WHERE table_region LIKE 'primary_%'
        ORDER BY physical_page, table_region, row_index, fiscal_year
    """
    for r in con.execute(q_primary):
        ws_primary.append([
            int(r["id"]),
            int(r["physical_page"]),
            str(r["table_region"]),
            str(r["row_label"]),
            str(r["fiscal_year"]),
            str(r["value_raw"]),
            float(r["value_num"]) if r["value_num"] is not None else None,
            str(r["unit"]),
            str(r["extraction_commit"]),
        ])

    # ---------------------------------------------- 2. Section D009 Cells Tab
    ws_d009 = wb.create_sheet(title="Production_Costs_p48_50")
    headers_d009 = [
        "Cell ID", "Physical Page", "Table Region", "Row Label", "Fiscal Year",
        "Col Index", "Column Label", "Value Raw", "Value Num", "Unit", "Extraction Commit"
    ]
    ws_d009.append(headers_d009)

    q_d009 = """
        SELECT id, physical_page, table_region, row_label, fiscal_year,
               column_index, column_label, value_raw, value_num, unit, extraction_commit
        FROM cells WHERE table_region NOT LIKE 'primary_%'
        ORDER BY physical_page, row_index, fiscal_year, column_index
    """
    for r in con.execute(q_d009):
        ws_d009.append([
            int(r["id"]),
            int(r["physical_page"]),
            str(r["table_region"]),
            str(r["row_label"]) if r["row_label"] is not None else "",
            str(r["fiscal_year"]) if r["fiscal_year"] is not None else "",
            int(r["column_index"]) if r["column_index"] is not None else None,
            str(r["column_label"]) if r["column_label"] is not None else "",
            str(r["value_raw"]),
            float(r["value_num"]) if r["value_num"] is not None else None,
            str(r["unit"]) if r["unit"] is not None else "",
            str(r["extraction_commit"]),
        ])

    # --------------------------------------------------------- 3. Conflicts Tab
    ws_conflicts = wb.create_sheet(title="Conflicts")
    conflict_headers = [
        "Conflict ID", "Check ID", "Kind", "Physical Page", "Ledger Question ID", "Status", "Description"
    ]
    ws_conflicts.append(conflict_headers)

    q_conflicts = """
        SELECT id, check_id, kind, physical_page, ledger_question_id, status, description
        FROM conflicts ORDER BY id
    """
    for r in con.execute(q_conflicts):
        ws_conflicts.append([
            int(r["id"]),
            str(r["check_id"]) if r["check_id"] is not None else "",
            str(r["kind"]),
            int(r["physical_page"]) if r["physical_page"] is not None else None,
            str(r["ledger_question_id"]) if r["ledger_question_id"] is not None else "",
            str(r["status"]),
            str(r["description"]),
        ])

    # --------------------------------------------------------- 4. Questions Tab
    ws_questions = wb.create_sheet(title="Questions")
    question_headers = [
        "Question ID", "Priority", "Status", "Target Document", "Target Section", "Question"
    ]
    ws_questions.append(question_headers)

    q_questions = """
        SELECT question_id, priority, status, target_document, target_section, question
        FROM questions ORDER BY question_id
    """
    for r in con.execute(q_questions):
        ws_questions.append([
            str(r["question_id"]),
            str(r["priority"]),
            str(r["status"]),
            str(r["target_document"]),
            str(r["target_section"]),
            str(r["question"]),
        ])

    con.close()

    # Style all sheets and auto-fit column widths
    for ws, fill in [(ws_primary, header_fill_green), (ws_d009, header_fill_blue), (ws_conflicts, header_fill_blue), (ws_questions, header_fill_blue)]:
        ws.views.sheetView[0].showGridLines = True
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 65)
            for cell in col[1:]:
                cell.font = regular_font
                cell.border = thin_border
                val = str(cell.value or "")
                if val.startswith("="):
                    raise ValueError(f"Covenant 2 violation: formula found in sheet {ws.title} at {cell.coordinate}: {val}")

    wb.save(XLSX_PATH)
    print(f"Generated consolidated values-only Excel view: {XLSX_PATH}")
    print(f"Sheets: {[s.title for s in wb.worksheets]}")
    print(f"Primary Statements cells: {ws_primary.max_row - 1}")
    print(f"Production & Costs cells: {ws_d009.max_row - 1}")
    print(f"Conflicts: {ws_conflicts.max_row - 1}")
    print(f"Questions: {ws_questions.max_row - 1}")


if __name__ == "__main__":
    build_excel_view()
